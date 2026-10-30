"""Build the pitch deck (docs/slides/pitch.pptx) on a light, minimal template.

    /tmp/pptx-venv/bin/python docs/slides/build_slides.py     # needs python-pptx, pillow

Import into Google Slides: File -> Import slides -> Upload -> pitch.pptx (speaker notes come along).
The last slides are empty layouts (content, two columns, image, statement) to duplicate.
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

DOCS = Path(__file__).resolve().parents[1]
OUT = DOCS / "slides" / "pitch.pptx"

# ---------------------------------------------------------------- template
FONT = "Lato"
ACCENT = RGBColor(0x1A, 0x5F, 0xB4)  # one blue
INK = RGBColor(0x14, 0x1A, 0x22)
MUTED = RGBColor(0x6B, 0x74, 0x80)
RULE = RGBColor(0xE4, 0xE7, 0xEB)
PANEL = RGBColor(0xF4, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC6, 0x28, 0x28)
GREEN = RGBColor(0x2E, 0x7D, 0x32)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
M = Inches(0.9)  # outer margin
CONTENT_W = W - 2 * M
_counter = {"n": 0}


def txt(slide, x, y, w, h, s, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(s if isinstance(s, list) else [s]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.line_spacing = align, spacing
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return box


def rect(slide, x, y, w, h, fill, rounded=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    if rounded:
        shp.adjustments[0] = 0.06
    return shp


def picture(slide, path, x, y, w, h):
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def chrome(slide, dark=False):
    """Footer: team tag left, slide number right. Same on every slide."""
    _counter["n"] += 1
    c = RGBColor(0xB8, 0xC4, 0xD6) if dark else MUTED
    txt(slide, M, H - Inches(0.55), Inches(8), Inches(0.3), "HalfSecuredTeam  ·  EHDS Track 3", 11, False, c)
    txt(slide, W - M - Inches(1), H - Inches(0.55), Inches(1), Inches(0.3), str(_counter["n"]), 11, False, c, PP_ALIGN.RIGHT)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def new(title=None, kicker=None):
    """Content slide: small kicker, large title, thin accent rule. Returns slide and the top of the body area."""
    s = prs.slides.add_slide(BLANK)
    y = Inches(0.75)
    if kicker:
        txt(s, M, y, CONTENT_W, Inches(0.3), kicker.upper(), 12, True, ACCENT)
        y += Inches(0.35)
    if title:
        txt(s, M, y, CONTENT_W, Inches(0.9), title, 34, True, INK)
        y += Inches(0.95)
        rect(s, M, y, Inches(0.7), Inches(0.05), ACCENT)
        y += Inches(0.45)
    chrome(s)
    return s, y


def slide_title(title, subtitle, tag):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, Inches(0.35), H, ACCENT)
    txt(s, M + Inches(0.3), Inches(2.3), Inches(11), Inches(1.9), title, 46, True, INK, spacing=1.05)
    txt(s, M + Inches(0.3), Inches(4.35), Inches(10.5), Inches(0.9), subtitle, 22, False, MUTED)
    txt(s, M + Inches(0.3), Inches(6.2), Inches(10), Inches(0.4), tag, 14, False, MUTED)
    chrome(s)
    return s


def slide_section(number, title, line=None):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, H, ACCENT)
    txt(s, M, Inches(2.2), Inches(4), Inches(1.2), number, 72, True, RGBColor(0xBF, 0xD4, 0xEE))
    txt(s, M, Inches(3.4), Inches(11), Inches(1.2), title, 44, True, WHITE)
    if line:
        txt(s, M, Inches(4.6), Inches(11), Inches(0.8), line, 20, False, RGBColor(0xDC, 0xE7, 0xF5))
    chrome(s, dark=True)
    return s


def slide_statement(sentence, small=None):
    s = prs.slides.add_slide(BLANK)
    rect(s, M, Inches(2.2), Inches(0.08), Inches(2.6), ACCENT)
    txt(s, M + Inches(0.45), Inches(2.2), Inches(11), Inches(2.6), sentence, 36, True, INK, anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)
    if small:
        txt(s, M + Inches(0.45), Inches(5.0), Inches(11), Inches(0.5), small, 16, False, MUTED)
    chrome(s)
    return s


def columns(slide, y, items, size=18):
    """Up to four columns: accent tick, bold heading, one short line each."""
    n = len(items)
    gap = Inches(0.5)
    cw = (CONTENT_W - gap * (n - 1)) / n
    y += Inches(0.6)
    for i, (head, body) in enumerate(items):
        x = M + i * (cw + gap)
        rect(slide, x, y, Inches(0.5), Inches(0.06), ACCENT)
        txt(slide, x, y + Inches(0.3), cw, Inches(1.1), head, 24, True, INK, spacing=1.1)
        txt(slide, x, y + Inches(1.45), cw, Inches(2.5), body, size, False, MUTED, spacing=1.35)


# ---------------------------------------------------------------- deck
s = slide_title("Health data access that explains itself",
                "An automated Health Data Access Body for the European Health Data Space",
                "HalfSecuredTeam  ·  Track 3  ·  Regulation (EU) 2025/327")
notes(s, "Hi everyone. We are HalfSecuredTeam and we built an automated Health Data Access Body for the European "
         "Health Data Space: an app that decides health-data requests the way the law requires, and shows exactly why.")

s, y = new("Weeks of waiting, two kinds of people", "The problem")
columns(s, y, [
    ("Researchers", "know the medicine and the question — not how to query a database"),
    ("Engineers", "write the query — but not the medical codes, nor whether the purpose is lawful"),
    ("The access body", "must check every request against the regulation and protect patients"),
])
txt(s, M, Inches(6.1), CONTENT_W, Inches(0.5), "We sit at the intersection.", 24, True, ACCENT)
notes(s, "Researchers know the medicine and what they want to find out, but not how to query a database. Engineers "
         "can write the query but do not know the codes or whether the purpose is allowed. And the access body in "
         "between has to check every request against the regulation and protect patients. Today that takes weeks. "
         "We sit at the intersection.")

s, y = new("Four ideas", "Our approach")
columns(s, y, [
    ("Law-grounded AI", "picks the article and quotes it"),
    ("Words → SQL by rules", "no query written per request"),
    ("Privacy inside the query", "small cells and opt-outs cannot leak"),
    ("Secure room, no exit", "records viewed as images, never downloaded"),
])
notes(s, "Four ideas. The AI does not decide on its own: it picks the article from the regulation and has to quote it. "
         "Definitions become SQL through fixed rules, so no one writes a query per request and the numbers are the "
         "same every run. Privacy lives inside the query: small cells and opt-outs cannot leak. And records never "
         "leave the secure room: viewed as images, never downloaded.")

s, y = new("One pipeline, three outcomes", "How it works")
picture(s, DOCS / "how-it-works.png", M, y, CONTENT_W, H - y - Inches(0.8))
notes(s, "A request in any EU language is translated once. The AI picks the rule from the law text and quotes it. "
         "Approved requests are parsed by rules into SQL; statistics go through the small-cell rule, record requests "
         "into a secure room with no download. Anyone flooding the service is blocked.")

s = slide_section("Demo", "Four requests and one attack", "R02 rejected  ·  R01 statistics  ·  R11 in German  ·  R12 secure room  ·  flood blocked")
notes(s, "Switch to the screen capture.\nR02: the insurer. Rejected, Art. 54(b), quoted.\nR01: diabetes by region, the table "
         "with the small-cell note.\nR11: the same in German.\nR12: records for a model; open the secure room; Download: 403.\n"
         "flood.py: 20 served, then 429.")

s, y = new("What works — and what does not yet", "Honest summary")
half = (CONTENT_W - Inches(0.6)) / 2
for i, (head, col, items) in enumerate([
    ("Works", GREEN, ["10 / 10 decisions and articles", "6 / 6 numbers, 0 privacy leaks", "10 / 10 consistent across 3 runs, 5.7 s",
                      "34 / 36 unseen requests, incl. German", "Secure room: images only, download 403"]),
    ("Not yet", RED, ["Sometimes too strict when a study mentions a prohibited topic", "Cited point within Art. 53 can vary",
                      "Staff submission list is a placeholder", "Screens can be photographed — true of every SPE"]),
]):
    x = M + i * (half + Inches(0.6))
    rect(s, x, y, Inches(0.5), Inches(0.05), col)
    txt(s, x, y + Inches(0.25), half, Inches(0.6), head, 22, True, col)
    txt(s, x, y + Inches(0.95), half, Inches(3.5), ["—  " + t for t in items], 16, False, INK, spacing=1.5)
txt(s, M, Inches(6.3), CONTENT_W, Inches(0.4), "github.com/MatthiasEbner2002/Track_3-HalfSecuredTeam", 14, True, ACCENT)
notes(s, "What works: all ten task decisions, all six numbers, no leaks, consistent across runs, and 34 of 36 requests "
         "we had never seen. What does not yet: the law check is sometimes too strict, and the staff list is a "
         "placeholder. Everything is in the repository. Thank you.")

s = slide_statement("Decisions in seconds instead of weeks — every one explained, identical on every run.",
                    "Thank you  ·  github.com/MatthiasEbner2002/Track_3-HalfSecuredTeam")

# ---------------------------------------------------------------- empty layouts to duplicate
s, y = new("Content slide title", "Kicker")
txt(s, M, y, CONTENT_W, Inches(3), ["Body text, 20 pt. One thought per line.", "Second line.", "Third line."], 20, False, INK, spacing=1.5)

s, y = new("Two columns", "Kicker")
columns(s, y, [("Left heading", "one short line"), ("Right heading", "one short line")])

s, y = new("Image slide", "Kicker")
rect(s, M, y, CONTENT_W, H - y - Inches(0.8), PANEL, rounded=True)
txt(s, M, y, CONTENT_W, H - y - Inches(0.8), "image", 16, False, MUTED, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

slide_statement("A single sentence that matters.", "optional small line")

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(f"wrote {OUT} ({len(prs.slides)} slides)")
